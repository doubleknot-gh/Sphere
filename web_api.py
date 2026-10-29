from pathlib import Path
import csv
import io
import zipfile

from fastapi import Depends, File, Form, HTTPException, UploadFile
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from database import get_db
from models import ChatMessage, Notice

MAX_UPLOAD = 10 * 1024 * 1024
MAX_ROWS = 20000


async def read_sheet(file):
    content = await file.read(MAX_UPLOAD + 1)
    if len(content) > MAX_UPLOAD:
        raise HTTPException(400, "10MB 이하의 파일을 선택해주세요.")
    name = (file.filename or "").lower()
    try:
        if name.endswith('.csv'):
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('cp949')
            source = csv.reader(io.StringIO(text))
            workbook = None
        elif name.endswith('.xlsx'):
            # Reject compressed workbooks that would expand beyond a reasonable limit.
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise HTTPException(400, "압축 해제 크기가 너무 큰 엑셀 파일입니다.")
            workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            source = workbook.active.iter_rows(values_only=True)
        else:
            raise HTTPException(400, "xlsx 또는 csv 파일만 지원합니다.")
        rows = []
        try:
            for row in source:
                cleaned = [str(int(v)) if isinstance(v, float) and v.is_integer()
                           else str(v).strip() if v is not None else '' for v in row]
                if any(cleaned):
                    rows.append(cleaned)
                if len(rows) > MAX_ROWS:
                    raise HTTPException(400, "한 번에 20,000행 이하로 업로드해주세요.")
                if len(cleaned) > 200:
                    raise HTTPException(400, "200열 이하의 파일을 사용해주세요.")
        finally:
            if workbook is not None:
                workbook.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "파일을 읽을 수 없습니다. 올바른 xlsx 또는 CSV 파일인지 확인해주세요.")
    if not rows:
        raise HTTPException(400, "파일에 데이터가 없습니다.")
    return rows


class ContentInput(BaseModel):
    content: str = Field(max_length=5000)


class MessageInput(BaseModel):
    message: str = Field(min_length=1, max_length=1000)


def install_web(app, current_user, current_admin, upload_handler):
    app.mount('/assets', StaticFiles(directory=Path(__file__).parent / 'web'), name='assets')

    @app.get('/health')
    def health():
        return {'status': 'ok'}

    @app.post('/admin/import-preview')
    async def preview(file: UploadFile = File(...), admin=Depends(current_admin)):
        rows = await read_sheet(file)
        return {'columns': rows[0], 'rows': rows[1:6], 'count': max(0, len(rows) - 1)}

    @app.post('/admin/import-members')
    async def import_members(file: UploadFile = File(...), name_column: int = Form(...),
                             sid_column: int = Form(...), club_column: int = Form(...),
                             has_header: bool = Form(True), session: Session = Depends(get_db),
                             admin=Depends(current_admin)):
        rows = await read_sheet(file)
        indices = [name_column, sid_column, club_column]
        if min(indices) < 0 or max(indices) >= len(rows[0]) or len(set(indices)) != 3:
            raise HTTPException(400, "이름·학번·소속을 서로 다른 열에 연결해주세요.")
        selected = rows[1:] if has_header else rows
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        for row in selected:
            if not any(row):
                continue
            if max(indices) >= len(row) or any(not row[i].strip() for i in indices):
                raise HTTPException(400, "이름·학번·소속이 비어 있는 행이 있습니다. 파일을 확인해주세요.")
            writer.writerow([row[i] for i in indices])
        if not selected:
            raise HTTPException(400, "등록할 회원이 없습니다.")
        normalized = UploadFile(filename='upload.csv', file=io.BytesIO(output.getvalue().encode('utf-8-sig')))
        return await upload_handler(file=normalized, db_session=session, admin=admin)

    @app.get('/notice')
    def notice(session: Session = Depends(get_db), user=Depends(current_user)):
        entry = session.get(Notice, 1)
        return {'content': entry.content if entry else ''}

    @app.put('/admin/notice')
    def save_notice(data: ContentInput, session: Session = Depends(get_db), admin=Depends(current_admin)):
        entry = session.get(Notice, 1)
        if entry is None:
            entry = Notice(id=1)
            session.add(entry)
        entry.content = data.content.strip()
        session.commit()
        return {'message': '공지사항을 저장했습니다.'}

    @app.get('/messages')
    def messages(session: Session = Depends(get_db), user=Depends(current_user)):
        entries = session.query(ChatMessage).order_by(ChatMessage.id.desc()).limit(100).all()
        return [{'id': m.id, 'name': m.name, 'club': m.club, 'role': m.role, 'message': m.message,
                 'created_at': m.created_at.isoformat() + 'Z',
                 'can_delete': user.role == 'admin' or user.student_id == m.student_id}
                for m in reversed(entries)]

    @app.post('/messages')
    def post_message(data: MessageInput, session: Session = Depends(get_db), user=Depends(current_user)):
        message = data.message.strip()
        if not message:
            raise HTTPException(400, '내용을 입력해주세요.')
        for word in ['시발', '씨발', '개새끼', '병신', '미친', '존나', '좆', '지랄', '염병', '느금마']:
            message = message.replace(word, '***')
        session.add(ChatMessage(student_id=user.student_id, name=user.name, club=user.club,
                                role=user.role, message=message))
        session.commit()
        return {'message': '메시지를 등록했습니다.'}

    @app.delete('/messages/{message_id}')
    def delete_message(message_id: int, session: Session = Depends(get_db), user=Depends(current_user)):
        entry = session.get(ChatMessage, message_id)
        if entry is None:
            raise HTTPException(404, '메시지를 찾을 수 없습니다.')
        if user.role != 'admin' and entry.student_id != user.student_id:
            raise HTTPException(403, '본인의 메시지만 삭제할 수 있습니다.')
        session.delete(entry)
        session.commit()
        return {'message': '메시지를 삭제했습니다.'}
